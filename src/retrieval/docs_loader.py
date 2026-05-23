import os
import re
from pypdf import PdfReader
""" from langchain_community.document_loaders import PyPDFLoader """
from docx import Document
from retrieval.chunking import chunk_text

""" Splits text into sections -----------------------------------------------------------------------------"""
def split_sections(text):
    pattern = r"(?=\bS?\.\d+\b|\b\d+\.\d+\b)"
    sections = re.split(pattern, text)
    return [s.strip() for s in sections if s.strip()]


""" Extracts text from a PDF file ---------------------------------------------------------------------------"""
def extract_text_from_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    """ reader = PyPDFLoader(file_path)  # Use PyPDFLoader for better text extraction """
    structured_blocks = []

    for page in reader.pages:
        extracted = page.extract_text()

        if not extracted:
            continue

        lines = extracted.split("\n")

        for line in lines:
            clean_line = line.strip()

            if not clean_line:
                continue

            # Check if there is a table: multiple columns separated by empty spaces
            parts = re.split(r"\s{2,}", clean_line)

            if len(parts) >= 3:
                # Considered a table´s row
                row_text = " | ".join(parts)
                structured_blocks.append(row_text)

            else:
                # normal text
                structured_blocks.append(clean_line)

    return "\n\n".join(structured_blocks)


""" Extracts text from a Word (.docx) file  ---------------------------------------------------------------------------"""
def extract_text_from_docx(file_path: str) -> str:
    doc = Document(file_path)
    extracted_blocks = []

    # Extract paragraph text
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            extracted_blocks.append(text)

    # Extract tables as structured rows
    for table in doc.tables:

        headers = [cell.text.strip() for cell in table.rows[0].cells]

        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]

            if not any(cells):
                continue

            structured_row = []

            for h, c in zip(headers, cells):
                if c:
                    structured_row.append(f"{h}: {c.strip()}")

            if structured_row:
                row_text = " | ".join(structured_row)
                extracted_blocks.append(row_text)

    # Convert list → structured string
    full_text = "\n\n".join(extracted_blocks)

    return full_text


""" Extracts text from a plain text file  ---------------------------------------------------------------------------"""
def extract_text_from_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()
    

""" Checks chunk content --------------------------------------------------------------------------------------------"""
def is_valid_chunk(text):
    words = text.split()    
    if len(words) < 6:         # If it is too short, False.
        return False    
    if len(set(words)) < 4:     # If there are less than 4 different words, False
        return False    
    if ":" in text and len(text.replace(":", "").strip()) < 10:  # If there are only symbols, False
        return False
    return True


""" Loads documents from a folder (PDF, DOCX, TXT), extracts text, applies chunking, and returns structured documents. -------------------- """
def load_documents_from_folder(folder_path: str):

    documents = []
    global_chunk_id = 0

    for filename in os.listdir(folder_path):

        file_path = os.path.join(folder_path, filename)
        full_text = ""

        # Extract text
        if filename.endswith(".pdf"):
            full_text = extract_text_from_pdf(file_path)

        elif filename.endswith(".docx"):
            full_text = extract_text_from_docx(file_path)

        elif filename.endswith(".txt"):
            full_text = extract_text_from_txt(file_path)

        else:
            continue

        # Skip if empty
        if not full_text.strip():
            continue

        # Dividir en secciones (sobre string completo)
        sections = split_sections(full_text)
        
        for section in sections:

            # Checks different sections
            match = re.search(r"\b(S?\.\d+|\d+\.\d+)\b", section)
            section_id = match.group(0) if match else "unknown"

            # Filter valid chunks
            if not is_valid_chunk(section):
                continue

            # chunking
            sub_chunks = chunk_text(section, chunk_size=300, overlap=80)
            
            for chunk in sub_chunks:
                documents.append({
                    "id": global_chunk_id,
                    "text": chunk,
                    "source": file_path,
                    "section": section_id,
                    "chunk_id": global_chunk_id
                })

                global_chunk_id += 1

    return documents


""" Exports chunked documents into a structured Word file ------------------------------------------------------------------"""
def export_documents_to_word(documents, output_path="output_chunks.docx"):
    doc = Document()
    doc.add_heading("Chunked Documents Analysis", level=0)

    current_source = None
    current_section = None

    for item in documents:

        source = item["source"]
        section = item["section"]
        text = item["text"]
        chunk_id = item["chunk_id"]

        if source != current_source:
            doc.add_page_break()
            doc.add_heading(f"Documento: {source}", level=1)
            current_source = source
            current_section = None  # reset sección

        if section != current_section:
            doc.add_heading(f"Sección: {section}", level=2)
            current_section = section

        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"Chunk {chunk_id}:\n")
        run.bold = True

        doc.add_paragraph(text)

    doc.save(output_path)